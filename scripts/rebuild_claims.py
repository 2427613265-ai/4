#!/usr/bin/env python3
"""Rebuild patent claims with i/i+1 numbered closed construction loops."""

from copy import deepcopy
from lxml import etree
from docx import Document
from docx.oxml.ns import qn

SRC = "/home/ubuntu/.cursor/projects/workspace/uploads/___3_PA26-1238-I-A-__-___________________________4966.docx"
OUT = "/workspace/docs/PA26-1238-I-A_权利要求书优化稿.docx"
CLAIMS_TXT = "/workspace/docs/权利要求书_优化稿.txt"

M = "{http://schemas.openxmlformats.org/officeDocument/2006/math}"
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
XML_SPACE = "{http://www.w3.org/XML/1998/namespace}space"

NSMAP_M = {"m": "http://schemas.openxmlformats.org/officeDocument/2006/math"}


def omml_plain(el):
    return "".join((t.text or "") for t in el.findall(".//m:t", NSMAP_M))


def make_run(text):
    r = etree.Element(W + "r")
    rPr = etree.SubElement(r, W + "rPr")
    rFonts = etree.SubElement(rPr, W + "rFonts")
    rFonts.set(qn("w:hint"), "eastAsia")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "楷体")
    rFonts.set(qn("w:cs"), "Times New Roman")
    sz = etree.SubElement(rPr, W + "sz")
    sz.set(qn("w:val"), "28")
    szCs = etree.SubElement(rPr, W + "szCs")
    szCs.set(qn("w:val"), "28")
    t = etree.SubElement(r, W + "t")
    if text[:1].isspace() or text[-1:].isspace():
        t.set(XML_SPACE, "preserve")
    t.text = text
    return r


def make_p(parts, *, center=False, empty=False):
    p = etree.Element(W + "p")
    pPr = etree.SubElement(p, W + "pPr")
    pStyle = etree.SubElement(pPr, W + "pStyle")
    pStyle.set(qn("w:val"), "4")
    widow = etree.SubElement(pPr, W + "widowControl")
    widow.set(qn("w:val"), "0")
    spacing = etree.SubElement(pPr, W + "spacing")
    if center:
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")
    else:
        spacing.set(qn("w:line"), "500")
        spacing.set(qn("w:lineRule"), "exact")
    if not center:
        ind = etree.SubElement(pPr, W + "ind")
        ind.set(qn("w:firstLine"), "560")
        ind.set(qn("w:firstLineChars"), "200")
    jc = etree.SubElement(pPr, W + "jc")
    jc.set(qn("w:val"), "center" if center else "both")
    rPr = etree.SubElement(pPr, W + "rPr")
    rFonts = etree.SubElement(rPr, W + "rFonts")
    rFonts.set(qn("w:hint"), "eastAsia")
    rFonts.set(qn("w:eastAsia"), "楷体")
    rFonts.set(qn("w:cs"), "Times New Roman")
    sz = etree.SubElement(rPr, W + "sz")
    sz.set(qn("w:val"), "28")
    szCs = etree.SubElement(rPr, W + "szCs")
    szCs.set(qn("w:val"), "28")
    if empty:
        return p
    for kind, val in parts:
        if kind == "t":
            p.append(make_run(val))
        elif kind == "m":
            p.append(deepcopy(val))
        else:
            raise ValueError(kind)
    return p


def make_sub(base, sub):
    omath = etree.Element(M + "oMath")
    sSub = etree.SubElement(omath, M + "sSub")
    e = etree.SubElement(sSub, M + "e")
    r = etree.SubElement(e, M + "r")
    rPr = etree.SubElement(r, M + "rPr")
    sty = etree.SubElement(rPr, M + "sty")
    sty.set(qn("m:val"), "p")
    wrPr = etree.SubElement(r, W + "rPr")
    rFonts = etree.SubElement(wrPr, W + "rFonts")
    rFonts.set(qn("w:hint"), "eastAsia")
    rFonts.set(qn("w:ascii"), "Times New Roman")
    rFonts.set(qn("w:hAnsi"), "Times New Roman")
    rFonts.set(qn("w:eastAsia"), "楷体")
    sz = etree.SubElement(wrPr, W + "sz")
    sz.set(qn("w:val"), "28")
    t = etree.SubElement(r, M + "t")
    t.text = base
    subel = etree.SubElement(sSub, M + "sub")
    r2 = etree.SubElement(subel, M + "r")
    t2 = etree.SubElement(r2, M + "t")
    t2.text = sub
    return omath


def replace_range(doc, start, end, new_els):
    """Replace paragraphs [start, end] inclusive with new_els. Indices are current."""
    paras = doc.paragraphs
    anchor = paras[start - 1]._element
    olds = [paras[i]._element for i in range(start, end + 1)]
    for el in reversed(new_els):
        anchor.addnext(el)
    for el in olds:
        parent = el.getparent()
        parent.remove(el)


def parts_to_text(parts):
    out = []
    for kind, val in parts:
        if kind == "t":
            out.append(val)
        else:
            s = omml_plain(val)
            out.append(s if s else "□")
    return "".join(out)


def main():
    doc = Document(SRC)
    omaths = doc.element.body.findall(".//m:oMath", NSMAP_M)
    # Claims-area math (first 24)
    H0 = deepcopy(omaths[0])
    H1 = deepcopy(omaths[1])
    eta = deepcopy(omaths[5])
    dj = deepcopy(omaths[7])
    rhs = deepcopy(omaths[8])  # δ+Σ_m(Vj,m·ηj,m)
    Vjm = deepcopy(omaths[10])
    hk = deepcopy(omaths[12])  # hk in Δhk
    Wenv = deepcopy(omaths[14])
    Veq = deepcopy(omaths[22])  # V0·(L/10000)/Wenv
    V0 = deepcopy(omaths[23])
    dt = make_sub("Δ", "t")
    hq = make_sub("H", "q")  # unused, keep H0/H1

    def P(*chunks, center=False):
        parts = []
        i = 0
        chunks = list(chunks)
        while i < len(chunks):
            item = chunks[i]
            if item == "t":
                parts.append(("t", chunks[i + 1]))
                i += 2
            elif item == "m":
                parts.append(("m", chunks[i + 1]))
                i += 2
            else:
                raise ValueError(item)
        return make_p(parts, center=center)

    E = make_p([], empty=True)

    # ---------- 权利要求书 ----------
    claims = []

    c1 = [
        P("t", "1.一种悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，包括以下步骤："),
        P(
            "t",
            "S1、将钢主梁节段依次吊装至悬索桥主缆的吊索上，并纵向连接成连续钢梁骨架；将预制混凝土桥面板铺设于所述连续钢梁骨架的顶面，相邻所述预制混凝土桥面板之间形成湿接缝；",
        ),
        P(
            "t",
            "S2、将所述湿接缝划分为吊索区湿接缝和非吊索区湿接缝；",
        ),
        P(
            "t",
            "S3、开始第一阶段浇筑，先浇筑预设跨中区段的所述非吊索区湿接缝；",
        ),
        P(
            "t",
            "S4、再由跨中向两侧索塔对第i浇筑批次的所述非吊索区湿接缝进行逐批次、对称且跳仓浇筑；其中，i为正整数，i的起始值为1；",
        ),
        P(
            "t",
            "S5、判断第i浇筑批次完成后是否仍存在尚未浇筑的非吊索区湿接缝，若是，将i+1赋值给i，返回步骤S4；若否，进入步骤S6；",
        ),
        P(
            "t",
            "S6、待所述第一阶段的湿接缝混凝土达到预设强度后，开始第二阶段浇筑，由跨中向两侧索塔对称浇筑第N浇筑批次的所述吊索区湿接缝；其中，N为正整数，N的起始值为1；",
        ),
        P(
            "t",
            "S7、判断第N浇筑批次完成后是否仍存在尚未浇筑的吊索区湿接缝，若是，将N+1赋值给N，返回步骤S6中所述第二阶段浇筑；若否，全桥湿接缝浇筑完成，形成组合截面。",
        ),
    ]
    claims.extend(c1)
    claims.append(E)

    c2def = [
        P(
            "t",
            "2.根据权利要求1所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述吊索区湿接缝为位于吊索的中心线两侧预设范围内的湿接缝；所述非吊索区湿接缝为除所述吊索区湿接缝以外的湿接缝；所述跨中区段为以跨中为中心对称延伸的区段；所述跳仓浇筑为沿纵桥向间隔至少一个吊索节段浇筑，所述吊索节段为相邻两吊索之间区段。",
        ),
    ]
    claims.extend(c2def)
    claims.append(E)

    c2 = [
        P(
            "t",
            "3.根据权利要求1所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述S3中先浇筑预设跨中区段的所述非吊索区湿接缝包括以下步骤：",
        ),
        P(
            "t",
            "S31、以主缆跨中为起点，在浇筑前获得初始主缆水平张力",
            "m",
            H0,
            "t",
            "，作为跨中区段扩展的张力基准；令q≥1，q为正整数，q的起始值为1；",
        ),
        P(
            "t",
            "S32、由跨中向两侧索塔方向，按距跨中由近及远的顺序，对称选取一组尚未浇筑的非吊索区湿接缝作为第q次浇筑对象并浇筑，获得浇筑后的当前主缆水平张力",
            "m",
            H1,
            "t",
            "；",
        ),
        P(
            "t",
            "S33、判断(",
            "m",
            H1,
            "t",
            "−",
            "m",
            H0,
            "t",
            ")/",
            "m",
            H0,
            "t",
            "是否小于或等于15%，若是，将q+1赋值给q，返回步骤S32，继续按距跨中由近及远的顺序对称浇筑下一组尚未浇筑的非吊索区湿接缝；若否，则判定预设跨中区段浇筑完成，并进入步骤S4。",
        ),
    ]
    claims.extend(c2)
    claims.append(E)

    c3 = [
        P(
            "t",
            "4.根据权利要求3所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述S4中由跨中向两侧索塔对第i浇筑批次的所述非吊索区湿接缝进行逐批次、对称且跳仓浇筑包括以下步骤：",
        ),
        P(
            "t",
            "S34、第i浇筑批次开始前，获取主缆跨中相对湿接缝浇筑前标高的当前竖向位移δ，向下为正；",
        ),
        P(
            "t",
            "S35、获得各节段单位浇筑方量对跨中竖向位移的影响系数",
            "m",
            eta,
            "t",
            "，所述影响系数",
            "m",
            eta,
            "t",
            "为浇筑1 m³湿接缝混凝土引起的跨中竖向位移响应，与δ同向为正，针对尚未浇筑的各对称节段对j，按",
            "m",
            dj,
            "t",
            "=",
            "m",
            rhs,
            "t",
            "计算预估跨中竖向位移",
            "m",
            dj,
            "t",
            "，其中，",
            "m",
            Vjm,
            "t",
            "为节段对j中各节段的计划浇筑方量；",
        ),
        P(
            "t",
            "S36、在沿纵向与已浇筑段间隔至少一个吊索节段的各对称节段对中，选取使",
            "m",
            dj,
            "t",
            "最接近0的对称节段对作为第i浇筑批次的浇筑对象并浇筑。",
        ),
    ]
    claims.extend(c3)
    claims.append(E)

    c4 = [
        P(
            "t",
            "5.根据权利要求1所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述S6中，在所述第一阶段的湿接缝混凝土达到预设强度、且主缆线形达到预设指标后开始所述第二阶段浇筑；所述预设指标为：所述第一阶段连续三个浇筑批次中，每一批次结束后的跨中竖向位移变化量Δv均满足|Δv|≤3mm，且每一批次结束后的索塔塔顶纵向位移变化量均≤5mm。",
        ),
    ]
    claims.extend(c4)
    claims.append(E)

    c5 = [
        P(
            "t",
            "6.根据权利要求1所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述S6中，在第N浇筑批次为所述第二阶段最末一批次之前，先执行以下步骤：",
        ),
        P(
            "t",
            "S71、复测主缆跨中、1/4跨和3/4跨的标高，得到各测点的实测标高，并计算各测点实测标高与该测点当前施工阶段目标标高之差Δ",
            "m",
            hk,
            "t",
            "，其中k为测点编号；",
        ),
        P(
            "t",
            "S72、取各测点中绝对值最大的Δ",
            "m",
            hk,
            "t",
            "作为控制偏差Δh；",
        ),
        P(
            "t",
            "S73、判断是否满足Δh>10mm，若是，则增大所述最末一批次的计划浇筑方量以加载下压主缆；判断是否满足Δh<-10mm，若是，则减小所述最末一批次的计划浇筑方量以减载抑制下沉；判断是否满足|Δh|≤10mm，若是，则保持所述最末一批次的计划浇筑方量不变；再按确定后的计划浇筑方量浇筑所述最末一批次。",
        ),
    ]
    claims.extend(c5)
    claims.append(E)

    c6 = [
        P(
            "t",
            "7.根据权利要求1所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，还包括以下步骤：",
        ),
        P(
            "t",
            "S81、在步骤S4开始前，沿所述主缆纵向布设多个竖向位移监测点；",
        ),
        P(
            "t",
            "S82、每完成第i浇筑批次或第N浇筑批次，采集各所述竖向位移监测点在当前施工阶段的累积竖向位移；",
        ),
        P(
            "t",
            "S83、按",
            "m",
            Wenv,
            "t",
            "=max(",
            "m",
            dt,
            "t",
            ")−min(",
            "m",
            dt,
            "t",
            ")计算主缆变形包络线宽度",
            "m",
            Wenv,
            "t",
            "，其中",
            "m",
            dt,
            "t",
            "为第t个所述竖向位移监测点在当前批次完成后的累积竖向位移，t=1,2,…,n，n为所述竖向位移监测点的总数；",
        ),
        P(
            "t",
            "S84、判断所述",
            "m",
            Wenv,
            "t",
            "是否超过预设阈值，若否，进入下一浇筑批次；若是，则将累积竖向位移最大或最小的监测点作为变形极值测点，将下一浇筑批次调整为先浇筑沿桥纵向与所述变形极值测点相邻、且浇筑后使所述",
            "m",
            Wenv,
            "t",
            "减小的尚未浇筑节段，和/或下调该批次的浇筑方量后，再进入下一浇筑批次。",
        ),
    ]
    claims.extend(c6)
    claims.append(E)

    c7 = [
        P(
            "t",
            "8.根据权利要求7所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述预设阈值为L/10000，L为所述主缆跨度；在所述",
            "m",
            Wenv,
            "t",
            "大于所述预设阈值的前提下，下调后的浇筑方量V按下式计算，且不小于设计方量的预设比例：",
        ),
        P("t", "V=", "m", Veq, center=True),
        P("t", "其中，", "m", V0, "t", "为该节段原计划设计方量。"),
    ]
    claims.extend(c7)
    claims.append(E)

    c8 = [
        P(
            "t",
            "9.根据权利要求1至8中任一项所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述湿接缝混凝土的强度等级不低于所述预制混凝土桥面板的强度等级。",
        ),
    ]
    claims.extend(c8)
    claims.append(E)

    c9 = [
        P(
            "t",
            "10.根据权利要求2所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述预设范围为吊索中心线两侧各1.25m。",
        ),
    ]
    claims.extend(c9)
    claims.append(E)

    c10 = [
        P(
            "t",
            "11.根据权利要求1所述的悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，所述第一阶段分为10-12个浇筑批次完成，每一批次的浇筑方量与设计方量之差的绝对值不超过设计方量的5%。",
        ),
    ]
    claims.extend(c10)

    # ---------- 发明内容（与权利要求对应） ----------
    invention = [
        P(
            "t",
            "为解决上述问题，本发明提供一种悬索桥钢混组合梁原位叠合浇筑施工方法，包括以下步骤：",
        ),
        P(
            "t",
            "S1、将钢主梁节段依次吊装至悬索桥主缆的吊索上，并纵向连接成连续钢梁骨架；将预制混凝土桥面板铺设于所述连续钢梁骨架顶面，相邻所述预制混凝土桥面板之间形成湿接缝；",
        ),
        P(
            "t",
            "S2、将所述湿接缝划分为吊索区湿接缝和其余非吊索区湿接缝；",
        ),
        P(
            "t",
            "S3、开始第一阶段浇筑，先浇筑预设跨中区段的所述非吊索区湿接缝；",
        ),
        P(
            "t",
            "S4、再由跨中向两侧索塔对第i浇筑批次的所述非吊索区湿接缝进行逐批次、对称且跳仓浇筑；其中，i为正整数，i的起始值为1；",
        ),
        P(
            "t",
            "S5、判断第i浇筑批次完成后是否仍存在尚未浇筑的非吊索区湿接缝，若是，将i+1赋值给i，返回步骤S4；若否，进入步骤S6；",
        ),
        P(
            "t",
            "S6、待所述第一阶段的湿接缝混凝土达到预设强度后，开始第二阶段浇筑，由跨中向两侧索塔对称浇筑第N浇筑批次的所述吊索区湿接缝；其中，N为正整数，N的起始值为1；",
        ),
        P(
            "t",
            "S7、判断第N浇筑批次完成后是否仍存在尚未浇筑的吊索区湿接缝，若是，将N+1赋值给N，返回步骤S6中所述第二阶段浇筑；若否，全桥湿接缝浇筑完成并形成组合截面。",
        ),
        P(
            "t",
            "进一步地，所述吊索区湿接缝为位于吊索中心线两侧预设范围内的湿接缝；所述非吊索区湿接缝为除所述吊索区湿接缝以外的湿接缝；所述跨中区段为以跨中为中心对称延伸的区段；所述跳仓浇筑为沿桥纵向间隔至少一个吊索节段浇筑，所述吊索节段为相邻两吊索之间区段。",
        ),
        P(
            "t",
            "进一步地，所述S3中预设跨中区段按以下步骤确定并浇筑：",
        ),
        P(
            "t",
            "S31、以主缆跨中为起点，在浇筑前获得初始主缆水平张力",
            "m",
            H0,
            "t",
            "，作为跨中区段扩展的张力基准；令q≥1，q为正整数，q的起始值为1；",
        ),
        P(
            "t",
            "S32、由跨中向两侧索塔方向，按距跨中由近及远的顺序，对称选取一组尚未浇筑的非吊索区湿接缝作为第q次浇筑对象并浇筑，获得浇筑后的当前主缆水平张力",
            "m",
            H1,
            "t",
            "；",
        ),
        P(
            "t",
            "S33、判断(",
            "m",
            H1,
            "t",
            "−",
            "m",
            H0,
            "t",
            ")/",
            "m",
            H0,
            "t",
            "是否小于或等于15%，若是，将q+1赋值给q，返回步骤S32，继续按距跨中由近及远的顺序对称浇筑下一组尚未浇筑的非吊索区湿接缝；若否，则判定预设跨中区段浇筑完成，并进入步骤S4。",
        ),
        P(
            "t",
            "进一步地，所述S4中由跨中向两侧索塔对第i浇筑批次的非吊索区湿接缝进行逐批次、对称且跳仓浇筑时，按照以下步骤进行：",
        ),
        P(
            "t",
            "S34、第i浇筑批次开始前，获取主缆跨中相对湿接缝浇筑前标高的当前竖向位移δ，向下为正；",
        ),
        P(
            "t",
            "S35、获得各节段单位浇筑方量对跨中竖向位移的影响系数",
            "m",
            eta,
            "t",
            "，所述影响系数",
            "m",
            eta,
            "t",
            "为浇筑1 m³湿接缝混凝土引起的跨中竖向位移响应，与δ同向为正，针对尚未浇筑的各对称节段对j，按",
            "m",
            dj,
            "t",
            "=",
            "m",
            rhs,
            "t",
            "计算预估跨中竖向位移",
            "m",
            dj,
            "t",
            "，其中",
            "m",
            Vjm,
            "t",
            "为节段对j中各节段的计划浇筑方量；",
        ),
        P(
            "t",
            "S36、在沿纵向与已浇筑段间隔至少一个吊索节段的各对称节段对中，选取使",
            "m",
            dj,
            "t",
            "最接近0的对称节段对作为第i浇筑批次的浇筑对象并浇筑。",
        ),
        P("t", "进一步地，所述预设范围为吊索中心线两侧各1.25m。"),
        P(
            "t",
            "进一步地，所述第一阶段分为10-12个浇筑批次完成，每一批次的浇筑方量与设计方量之差的绝对值不超过设计方量的5%。",
        ),
        P(
            "t",
            "进一步地，所述S6中，在所述第一阶段的湿接缝混凝土达到预设强度、且主缆线形达到预设指标后开始所述第二阶段浇筑；所述预设指标为：所述第一阶段连续三个浇筑批次中，每一批次结束后的跨中竖向位移变化量Δv均满足|Δv|≤3mm，且每一批次结束后的索塔塔顶纵向位移变化量均≤5mm。",
        ),
        P(
            "t",
            "进一步地，所述S6中，在第N浇筑批次为所述第二阶段最末一批次之前，复测主缆跨中、1/4跨和3/4跨的标高，得到各测点的实测标高，并计算各测点实测标高与该测点当前施工阶段目标标高之差Δ",
            "m",
            hk,
            "t",
            "，其中k为测点编号；取各测点中绝对值最大的Δ",
            "m",
            hk,
            "t",
            "作为控制偏差Δh；判断是否满足Δh>10mm，若是，则增大所述最末一批次的计划浇筑方量以加载下压主缆；判断是否满足Δh<-10mm，若是，则减小所述最末一批次的计划浇筑方量以减载抑制下沉；判断是否满足|Δh|≤10mm，若是，则保持所述最末一批次的计划浇筑方量不变；再按确定后的计划浇筑方量浇筑所述最末一批次。",
        ),
        P("t", "进一步地，还包括以下步骤："),
        P("t", "S81、在步骤S4开始前，沿所述主缆纵向布设多个竖向位移监测点；"),
        P(
            "t",
            "S82、每完成第i浇筑批次或第N浇筑批次，采集各所述竖向位移监测点在当前施工阶段的累积竖向位移；",
        ),
        P(
            "t",
            "S83、按",
            "m",
            Wenv,
            "t",
            "=max(",
            "m",
            dt,
            "t",
            ")−min(",
            "m",
            dt,
            "t",
            ")计算主缆变形包络线宽度",
            "m",
            Wenv,
            "t",
            "，其中",
            "m",
            dt,
            "t",
            "为第t个所述竖向位移监测点在当前批次完成后的累积竖向位移，t=1,2,…,n，n为所述竖向位移监测点的总数；",
        ),
        P(
            "t",
            "S84、判断所述",
            "m",
            Wenv,
            "t",
            "是否超过预设阈值，若否，进入下一浇筑批次；若是，则将累积竖向位移最大或最小的监测点作为变形极值测点，将下一浇筑批次调整为先浇筑沿桥纵向与所述变形极值测点相邻、且浇筑后使所述",
            "m",
            Wenv,
            "t",
            "减小的尚未浇筑节段，和/或下调该批次的浇筑方量后，再进入下一浇筑批次。",
        ),
        P(
            "t",
            "进一步地，所述预设阈值为L/10000，L为所述主缆跨度；在所述",
            "m",
            Wenv,
            "t",
            "大于所述预设阈值的前提下，下调后的浇筑方量V按下式计算，且不小于设计方量的预设比例：",
        ),
        P("t", "V=", "m", Veq, center=True),
        P("t", "其中，", "m", V0, "t", "为该节段原计划设计方量。"),
        P(
            "t",
            "进一步地，所述湿接缝混凝土的强度等级不低于所述预制混凝土桥面板的强度等级。",
        ),
    ]

    embodiment_intro = [
        P(
            "t",
            "在本实施例中，如图1所示，一种悬索桥钢混组合梁原位叠合浇筑施工方法，包括以下步骤：",
        ),
        P(
            "t",
            "S1、将钢主梁节段依次吊装至悬索桥主缆的吊索上，并纵向连接成连续钢梁骨架；将预制混凝土桥面板铺设于所述连续钢梁骨架顶面，相邻所述预制混凝土桥面板之间形成湿接缝；",
        ),
        P(
            "t",
            "S2、将所述湿接缝划分为吊索区湿接缝和其余非吊索区湿接缝；",
        ),
        P(
            "t",
            "S3、开始第一阶段浇筑，先浇筑预设跨中区段的所述非吊索区湿接缝；",
        ),
        P(
            "t",
            "S4、再由跨中向两侧索塔对第i浇筑批次的所述非吊索区湿接缝进行逐批次、对称且跳仓浇筑；其中，i为正整数，i的起始值为1；",
        ),
        P(
            "t",
            "S5、判断第i浇筑批次完成后是否仍存在尚未浇筑的非吊索区湿接缝，若是，将i+1赋值给i，返回步骤S4；若否，进入步骤S6；",
        ),
        P(
            "t",
            "S6、待所述第一阶段的湿接缝混凝土达到预设强度后，开始第二阶段浇筑，由跨中向两侧索塔对称浇筑第N浇筑批次的所述吊索区湿接缝；其中，N为正整数，N的起始值为1；",
        ),
        P(
            "t",
            "S7、判断第N浇筑批次完成后是否仍存在尚未浇筑的吊索区湿接缝，若是，将N+1赋值给N，返回步骤S6中所述第二阶段浇筑；若否，全桥湿接缝浇筑完成并形成组合截面。",
        ),
    ]

    s31_block = [
        P(
            "t",
            "S31、以主缆跨中为起点，在浇筑前获得初始主缆水平张力",
            "m",
            H0,
            "t",
            "，作为跨中区段扩展的张力基准；令q≥1，q为正整数，q的起始值为1；",
        ),
        P(
            "t",
            "S32、由跨中向两侧索塔方向，按距跨中由近及远的顺序，对称选取一组尚未浇筑的非吊索区湿接缝作为第q次浇筑对象并浇筑，获得浇筑后的当前主缆水平张力",
            "m",
            H1,
            "t",
            "；",
        ),
        P(
            "t",
            "S33、判断(",
            "m",
            H1,
            "t",
            "−",
            "m",
            H0,
            "t",
            ")/",
            "m",
            H0,
            "t",
            "是否小于或等于15%，若是，将q+1赋值给q，返回步骤S32，继续按距跨中由近及远的顺序对称浇筑下一组尚未浇筑的非吊索区湿接缝；若否，则判定预设跨中区段浇筑完成，并进入步骤S4。",
        ),
    ]

    s34_block = [
        P(
            "t",
            "S34、第i浇筑批次开始前，获取主缆跨中相对湿接缝浇筑前标高的当前竖向位移δ，向下为正；",
        ),
        P(
            "t",
            "S35、获得各节段单位浇筑方量对跨中竖向位移的影响系数",
            "m",
            eta,
            "t",
            "，所述影响系数",
            "m",
            eta,
            "t",
            "为浇筑1 m³湿接缝混凝土引起的跨中竖向位移响应，与δ同向为正，针对尚未浇筑的各对称节段对j，按",
            "m",
            dj,
            "t",
            "=",
            "m",
            rhs,
            "t",
            "计算预估跨中竖向位移",
            "m",
            dj,
            "t",
            "，其中",
            "m",
            Vjm,
            "t",
            "为节段对j中各节段的计划浇筑方量；",
        ),
        P(
            "t",
            "S36、在沿纵向与已浇筑段间隔至少一个吊索节段的各对称节段对中，选取使",
            "m",
            dj,
            "t",
            "最接近0的对称节段对作为第i浇筑批次的浇筑对象并浇筑。",
        ),
    ]

    def ptext(p):
        return p.text or ""

    def find_para(prefix, start=0):
        for i, p in enumerate(doc.paragraphs):
            if i < start:
                continue
            if ptext(p).startswith(prefix):
                return i
        raise KeyError(prefix)

    def find_last(prefix):
        found = None
        for i, p in enumerate(doc.paragraphs):
            if ptext(p).startswith(prefix):
                found = i
        if found is None:
            raise KeyError(prefix)
        return found

    def find_contains(substr, start=0):
        for i, p in enumerate(doc.paragraphs):
            if i < start:
                continue
            if substr in ptext(p):
                return i
        raise KeyError(substr)

    # Embodiment S34-S36 (last occurrence is 具体实施方式)
    i34 = find_last("S34、每一浇筑批次开始前")
    i36 = find_para(
        "S36、在沿纵向与已浇筑段间隔至少一个吊索节段的各对称节段对中，选取使",
        i34,
    )
    replace_range(doc, i34, i36, s34_block)

    # Embodiment S31-S33 (last occurrence is 具体实施方式)
    last31 = find_last(
        "S31、以主缆跨中为起点，向两侧索塔方向对称选取尚未浇筑的非吊索区湿接缝，在浇筑前获得初始主缆水平张力"
    )
    i33 = find_para("S33、判断", last31)
    replace_range(doc, last31, i33, s31_block)

    # Embodiment intro S1-S4
    intro = find_para("在本实施例中，如图1所示")
    s4e = find_para("S4、判断所述第一阶段的湿接缝混凝土是否达到预设强度", intro)
    replace_range(doc, intro, s4e, embodiment_intro)

    # 发明内容 method listing
    inv = find_para(
        "为解决上述问题，本发明提供一种悬索桥钢混组合梁原位叠合浇筑施工方法，包括以下步骤："
    )
    inv_end = find_para(
        "进一步地，所述湿接缝混凝土的强度等级不低于所述预制混凝土桥面板的强度等级。",
        inv,
    )
    replace_range(doc, inv, inv_end, invention)

    # 权利要求书
    cstart = find_para(
        "1.一种悬索桥钢混组合梁原位叠合浇筑施工方法，其特征在于，包括以下步骤："
    )
    cend = find_contains("所述第一阶段分为10")
    replace_range(doc, cstart, cend, claims)

    # Align leftover embodiment explanations with the closed q-loop (H0 obtained once in S31).
    replacements = [
        (
            "所述初始主缆水平张力H₀是指对当前选取的非吊索区湿接缝实施浇筑之前、对应施工阶段下主缆的水平张力",
            "所述初始主缆水平张力H₀是指预设跨中区段浇筑开始前、由步骤S31获得的主缆水平张力，并在q循环内保持不变",
        ),
        (
            "选取后、浇筑前先获得H₀，作为本轮跨中扩展的张力基准；浇筑完成后，将本次湿接缝混凝土自重计入当前施工阶段荷载，再获得H₁，并计算(H₁−H₀)/H₀。",
            "在步骤S31于浇筑前一次获得H₀，作为本轮跨中扩展的张力基准；此后由跨中向两侧、按距跨中由近及远的顺序，将一组尚未浇筑的非吊索区湿接缝作为第q次浇筑对象并浇筑，浇筑完成后获得H₁，并计算(H₁−H₀)/H₀。若(H₁−H₀)/H₀≤15%，将q+1赋值给q并返回步骤S32，H₀保持不变，继续浇筑下一组；若(H₁−H₀)/H₀>15%，判定跨中区段浇筑完成，进入步骤S4。",
        ),
        (
            "所述影响系数是指节段对j中第m个节段在单位浇筑荷载作用下引起的跨中竖向位移响应",
            "所述影响系数是指节段对j中第m个节段浇筑1 m³湿接缝混凝土时引起的跨中竖向位移响应",
        ),
        (
            "对相应节段施加单位浇筑荷载或等效单位湿接缝自重，读取跨中竖向位移响应得到",
            "对相应节段施加1 m³湿接缝混凝土对应的等效自重，读取跨中竖向位移响应得到",
        ),
    ]
    for p in doc.paragraphs:
        t = p.text or ""
        for old, new in replacements:
            if old in t and p.runs:
                full = t.replace(old, new)
                p.runs[0].text = full
                for r in p.runs[1:]:
                    r.text = ""

    doc.save(OUT)

    # Write plaintext claims for review
    lines = []
    for el in claims:
        texts = []
        for node in el:
            tag = node.tag
            if tag == W + "r":
                ts = node.findall(W + "t")
                texts.append("".join(t.text or "" for t in ts))
            elif tag.endswith("oMath") or tag == M + "oMath":
                texts.append(omml_plain(node))
        line = "".join(texts)
        lines.append(line)
    with open(CLAIMS_TXT, "w", encoding="utf-8") as f:
        f.write("\n".join(lines))
        f.write("\n")
    print("saved", OUT)
    print("saved", CLAIMS_TXT)
    print("--- claims text ---")
    print("\n".join(lines))


if __name__ == "__main__":
    main()
