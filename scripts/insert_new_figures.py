#!/usr/bin/env python3
"""Insert Fig.2 (zone colors) and Fig.3 (stage-1 pour sequence) into the application."""

from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt
from docx.text.paragraph import Paragraph

DOC = Path("/workspace/docs/PA26-1238-I-A_权利要求书优化稿.docx")
FIG2 = Path("/workspace/docs/figures/图2_湿接缝分区.png")
FIG3 = Path("/workspace/docs/figures/图3_第一阶段浇筑次序.png")
DESC_TXT = Path("/workspace/docs/附图说明_优化稿.txt")
EMB_TXT = Path("/workspace/docs/具体实施方式_优化稿.txt")

FIG2_TEXT = (
    "如图2所示，预制混凝土桥面板铺设于连续钢梁骨架顶面后，相邻预制板之间形成纵、横湿接缝。"
    "图2中红色填充的湿接缝为吊索区湿接缝，位于吊索中心线两侧预设范围内、对应吊点附近的接缝区段；"
    "蓝色填充的湿接缝为非吊索区湿接缝，即除吊索区湿接缝以外的其余湿接缝。"
    "第一阶段仅浇筑图2中蓝色所示的非吊索区湿接缝；待其达到预设强度后，第二阶段再浇筑图2中红色所示的吊索区湿接缝，"
    "以推迟吊点附近组合约束的形成。"
)

FIG3_TEXT = (
    "如图3所示，第一阶段非吊索区湿接缝沿主缆跨径按第1次至第10次共十个浇筑批次组织施工，并关于跨中左右对称。"
    "图3中D1～D65为沿纵桥向依次布置的吊索编号。"
    "第1次对称布置于跨中附近，对应于步骤S3优先浇筑的预设跨中区段；"
    "第2次至第5次由跨中向两侧索塔方向依次外扩，且与相邻已浇批次沿纵桥向间隔至少一个吊索节段；"
    "第6次至第10次再由跨中向两侧索塔方向回填前述间隔仓位，同样保持左右对称。"
    "图3所示各次浇筑对象均为非吊索区湿接缝，对应于图2中的蓝色填充；吊索区湿接缝（图2中红色填充）不在第一阶段浇筑。"
    "该次序体现了跳仓浇筑：相邻浇筑批次沿纵桥向不连续推进，从而避免局部刚度突变，"
    "并在不设置外部等效配重的条件下，形成“先跨中建立刚度、后两端对称跳仓对消变形、再回填间隔仓”的自平衡过程。"
)

CAP2 = "图2为湿接缝分区示意图，图中红色填充表示吊索区湿接缝，蓝色填充表示非吊索区湿接缝。"
CAP3 = "图3为第一阶段非吊索区湿接缝浇筑次序示意图。"
CAP3_MARKS = "图3中，D1～D65为沿纵桥向依次布置的吊索编号，第1次～第10次表示第一阶段非吊索区湿接缝的浇筑批次。"


def find_para(doc, predicate):
    hits = [p for p in doc.paragraphs if predicate(p)]
    if len(hits) != 1:
        raise RuntimeError(f"expected 1 hit, got {len(hits)} for {predicate}")
    return hits[0]


def set_text_keep_format(p, text):
    if p.runs:
        p.runs[0].text = text
        for r in p.runs[1:]:
            r.text = ""
    else:
        p.add_run(text)


def clone_empty_after(paragraph):
    new_el = deepcopy(paragraph._p)
    pPr = new_el.find(qn("w:pPr"))
    for child in list(new_el):
        if child is not pPr:
            new_el.remove(child)
    if pPr is not None and pPr.find(qn("w14:paraId")) is not None:
        pPr.remove(pPr.find(qn("w14:paraId")))
    paragraph._p.addnext(new_el)
    return Paragraph(new_el, paragraph._parent)


def clone_text_after(paragraph, text):
    new_p = clone_empty_after(paragraph)
    run = new_p.add_run(text)
    if paragraph.runs:
        src = paragraph.runs[0]
        run.font.size = src.font.size
        run.font.name = src.font.name
        rPr = run._element.get_or_add_rPr()
        src_rPr = src._element.find(qn("w:rPr"))
        if src_rPr is not None:
            for child in list(src_rPr):
                rPr.append(deepcopy(child))
    return new_p


def replace_in_runs(p, old, new):
    full = p.text or ""
    if old not in full:
        raise RuntimeError(f"not found: {old!r} in {full[:80]!r}")
    if p.runs:
        p.runs[0].text = full.replace(old, new)
        for r in p.runs[1:]:
            r.text = ""


def main():
    doc = Document(str(DOC))

    p_fig1 = find_para(
        doc,
        lambda p: (p.text or "").strip()
        == "图1为本发明悬索桥钢混组合梁原位叠合浇筑施工方法的整体流程图。",
    )
    # insert after fig1 caption: fig2, fig3, fig3 marks (addnext reverse order)
    clone_text_after(p_fig1, CAP3_MARKS)
    clone_text_after(p_fig1, CAP3)
    clone_text_after(p_fig1, CAP2)

    p_start = find_para(
        doc,
        lambda p: (p.text or "").startswith("在本实施例中，如图1所示，一种悬索桥钢混组合梁"),
    )
    replace_in_runs(p_start, "如图1所示", "如图1至图3所示")

    p_zone = find_para(
        doc,
        lambda p: (p.text or "").startswith("在浇筑前，将全桥湿接缝划分为吊索区湿接缝与非吊索区湿接缝。"),
    )
    clone_text_after(p_zone, FIG2_TEXT)

    p_stage1 = find_para(
        doc,
        lambda p: (p.text or "").startswith("第一阶段仅浇筑非吊索区湿接缝，且不依赖水箱、砂袋等外部等效配重。"),
    )
    clone_text_after(p_stage1, FIG3_TEXT)

    p_stage2 = find_para(
        doc,
        lambda p: (p.text or "").startswith("当判断第一阶段湿接缝混凝土已达到预设强度后，进入第二阶段"),
    )
    replace_in_runs(
        p_stage2,
        "进入第二阶段，由跨中向两侧索塔对称、逐批次浇筑吊索区湿接缝。",
        "进入第二阶段，由跨中向两侧索塔对称、逐批次浇筑图2中红色所示的吊索区湿接缝。",
    )

    p_range = find_para(
        doc,
        lambda p: (p.text or "").startswith("具体而言，由此形成的吊索区总宽约2.5m"),
    )
    replace_in_runs(
        p_range,
        "具体而言，由此形成的吊索区总宽约2.5m，对应吊点附近需要滞后闭合的湿接缝范围。",
        "具体而言，由此形成的吊索区总宽约2.5m，对应吊点附近需要滞后闭合的湿接缝范围，即图2中红色填充所示区域。",
    )

    p_batches = find_para(
        doc,
        lambda p: (p.text or "").startswith("具体而言，第一阶段批次划分过少"),
    )
    replace_in_runs(
        p_batches,
        "10–12批与主跨数百米至千米级、吊索节段数十个量级的工程相匹配，可使荷载增加→张力提升→后续变形抑制的过程阶梯推进。",
        "本实施例图3给出了第一阶段分为10个浇筑批次、左右对称施工的示例；10–12批与主跨数百米至千米级、吊索节段数十个量级的工程相匹配，可使荷载增加→张力提升→后续变形抑制的过程阶梯推进。",
    )

    p_cap = find_para(doc, lambda p: (p.text or "").strip() == "图1")
    # after 图1 caption: image2, 图2, image3, 图3  (addnext reverse)
    def add_caption_and_picture(after_para, caption, image_path, width_cm):
        cap = clone_empty_after(after_para)
        cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = cap.add_run(caption)
        r.font.size = Pt(14)
        r.font.name = "Times New Roman"
        rPr = r._element.get_or_add_rPr()
        rFonts = rPr.find(qn("w:rFonts"))
        if rFonts is None:
            rFonts = r._element.makeelement(qn("w:rFonts"), {})
            rPr.insert(0, rFonts)
        rFonts.set(qn("w:ascii"), "Times New Roman")
        rFonts.set(qn("w:hAnsi"), "Times New Roman")
        rFonts.set(qn("w:eastAsia"), "楷体")
        img_p = clone_empty_after(after_para)
        img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        img_p.add_run().add_picture(str(image_path), width=Cm(width_cm))
        return cap

    add_caption_and_picture(p_cap, "图3", FIG3, 16.0)
    add_caption_and_picture(p_cap, "图2", FIG2, 9.0)

    doc.save(str(DOC))

    # dump 附图说明 + 具体实施方式 for review
    paras = [p.text.strip() for p in Document(str(DOC)).paragraphs]
    start = paras.index("附图说明")
    end = paras.index("说明书   附图".replace(" ", "")) if False else None
    # find 说明书附图 heading
    fig_head = None
    for i, t in enumerate(paras):
        if "说明书" in t and "附图" in t and i > start:
            fig_head = i
            break
    body = paras[start:fig_head]
    DESC_TXT.write_text("\n".join([t for t in body if t][:8]) + "\n", encoding="utf-8")
    # rewrite 附图说明 from canonical text
    DESC_TXT.write_text(
        "附图说明\n\n"
        "为了更清楚地说明本发明实施例或现有技术中的技术方案，下面将对实施例或现有技术描述中所需要使用的附图作简单地介绍，"
        "显而易见地，下面描述中的附图仅仅是本发明的一些实施例，对于本领域普通技术人员来讲，在不付出创造性劳动的前提下，"
        "还可以根据这些附图示出的结构获得其他的附图。\n\n"
        f"{CAP2.replace('图2为', '图1为本发明悬索桥钢混组合梁原位叠合浇筑施工方法的整体流程图。\n图2为', 1)}\n"
        f"{CAP3}\n"
        f"{CAP3_MARKS}\n".replace(
            "图2为湿接缝分区示意图，图中红色填充表示吊索区湿接缝，蓝色填充表示非吊索区湿接缝。\n图1为",
            "图1为",
        ),
        encoding="utf-8",
    )
    # simpler write
    DESC_TXT.write_text(
        "附图说明\n\n"
        "为了更清楚地说明本发明实施例或现有技术中的技术方案，下面将对实施例或现有技术描述中所需要使用的附图作简单地介绍，"
        "显而易见地，下面描述中的附图仅仅是本发明的一些实施例，对于本领域普通技术人员来讲，在不付出创造性劳动的前提下，"
        "还可以根据这些附图示出的结构获得其他的附图。\n\n"
        "图1为本发明悬索桥钢混组合梁原位叠合浇筑施工方法的整体流程图。\n"
        f"{CAP2}\n"
        f"{CAP3}\n"
        f"{CAP3_MARKS}\n",
        encoding="utf-8",
    )
    print("saved", DOC)
    print("saved", DESC_TXT)
    # verify
    d2 = Document(str(DOC))
    print("inline_shapes", len(d2.inline_shapes))
    for i, p in enumerate(d2.paragraphs, 1):
        t = (p.text or "").strip()
        if t.startswith("图") or t.startswith("如图") or "如图1至图3" in t or "红色填充" in t:
            print(f"[{i}] {t[:120]}")


if __name__ == "__main__":
    main()
