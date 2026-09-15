# -*- coding: utf-8 -*-
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement


def set_run_font(run, cn="宋体", en="Times New Roman", size=12, bold=False, color=None):
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = en
    if color is not None:
        run.font.color.rgb = color
    r = run._element
    rPr = r.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), en)
    rFonts.set(qn("w:hAnsi"), en)
    rFonts.set(qn("w:eastAsia"), cn)
    rFonts.set(qn("w:cs"), en)


def add_para(doc, text, *, cn="宋体", size=12, bold=False, first_line=True, space_after=8, align="justify"):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if first_line:
        pf.first_line_indent = Cm(0.74)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    elif align == "left":
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    run = p.add_run(text)
    set_run_font(run, cn=cn, size=size, bold=bold)
    return p


def add_mixed_para(doc, parts, *, first_line=True, space_after=8, align="justify", size=12):
    """parts: list of (text, bold)"""
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.5
    if first_line:
        pf.first_line_indent = Cm(0.74)
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.first_line_indent = Cm(0)
    else:
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for text, bold in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold)
    return p


def add_heading_cn(doc, text, size=16):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(18)
    pf.space_after = Pt(12)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(text)
    set_run_font(run, cn="黑体", size=size, bold=True)
    return p


def add_section(doc, text, size=14):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(14)
    pf.space_after = Pt(8)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, cn="黑体", size=size, bold=True)
    return p


def add_subsection(doc, text):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(10)
    pf.space_after = Pt(6)
    pf.line_spacing = 1.5
    pf.first_line_indent = Cm(0)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, cn="黑体", size=12, bold=True)
    return p


def set_cell_shading(cell, fill):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")
    tcPr.append(shd)


def set_cell_text(cell, text, *, bold=False, size=10.5, center=False):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(2)
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    set_run_font(run, cn="宋体", size=size, bold=bold)


doc = Document()
for section in doc.sections:
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.8)
    section.right_margin = Cm(2.6)

add_heading_cn(doc, "实用新型专利检索报告（分析部分）", size=18)
add_para(
    doc,
    "案号：PA26-0865-U-A　　发明创造名称：一种用于截骨矫形的骨假体　　申请人：湘雅二医院",
    first_line=False,
    align="center",
    size=10.5,
    space_after=4,
)
add_para(
    doc,
    "对比文件1：CN116965979A《腿骨假体组件及膝关节假体》　　对比文件2：CN209253032U《一种用于胫骨近端高位截骨的矫正装置》",
    first_line=False,
    align="center",
    size=10.5,
    space_after=16,
)

add_section(doc, "一、本案独立权利要求所解决的技术问题及核心技术手段")
add_para(
    doc,
    "本案独立权利要求请求保护一种用于截骨矫形的骨假体，其针对长管状骨（如股骨、胫骨）因力线异常或成角畸形而行截骨矫形时，两骨段主要依赖平面或楔形断面贴合对位、断端相对自由度有限，难以在术中对涉及多个解剖学平面的成角偏差进行全向、精细调整，且矫正后的力线在恢复期不易稳定维持的问题，采用如下核心技术手段：第一假体与第二假体分别包括髓内柄及连接部；两连接部中一者设外球面、另一者设相适配的内球面，构成可相对转动的球面副，第一假体与第二假体经该球面副相互连接；第一髓内柄与第二髓内柄自球面副向相反的两侧延伸，并分别插入截骨后所形成的两个骨段的髓腔，以调整、维持力线。",
)

add_section(doc, "二、对比文件1的分析")
add_subsection(doc, "（一）对比文件1公开的技术方案")
add_para(
    doc,
    "对比文件1（CN116965979A）公开一种腿骨假体组件及膝关节假体，属于膝关节置换领域。该腿骨假体组件包括骨连接件、髓内柄和固定部件：骨连接件配置为股骨髁假体或胫骨托假体，用于连接至膝关节置换所准备的股骨远端或胫骨近端；髓内柄仅一根，其连接端与骨连接件之间以球面副可相对转动地配合，插入端插入同一骨骼的髓腔；固定部件与髓内柄之间亦以球面副配合，并将髓内柄连接、锁紧至骨连接件。其技术目的在于：在膝关节置换中根据患者股骨前弓角、外翻角或胫骨后倾角，调节该单根髓内柄相对已确定关节面假体的倾斜角度，使柄轴与患者解剖轴对齐，降低延长杆与皮质骨接触风险，减少库存规格。",
)

add_subsection(doc, "（二）与本案的技术构思及功能对比")
add_para(
    doc,
    "对比文件1采用“单柄—关节面球面适配”的被动匹配式调节：预先将骨连接件（股骨髁假体或胫骨托）固定于膝关节置换的外科准备面，仅将一根髓内柄插入同一骨骼的髓腔，按照球面副允许的万向转动，调整该髓内柄相对已确定关节面的倾角，未达匹配角度则继续转动，达到与解剖轴对齐后再由固定部件锁紧。其功能终点是使膝关节置换假体的髓内柄适配既有髓腔形态，球面副解决的是“假体关节面已定、单侧髓内柄安装角度可调”的装配问题，并不连接截骨后彼此分离的两个骨段，也不承担两骨段之间力线的建立与维持。",
)
add_para(
    doc,
    "本专利采用“双柄对向植入＋球面副枢纽全向矫形”的主动规划式调节：第一假体与第二假体分别插入长管状骨截骨后形成的两个骨段髓腔，两连接部构成球面副，作为跨越截骨线的中间枢纽；术中以球面副为中心，对两骨段断端相对角度作全向、连续、精细调整，调整到位后锁定，并借助两髓内柄在髓腔内的植入关系维持已矫正的力线。其功能终点是截骨矫形中“两骨段相对力线的建立、调整与术后维持”。",
)
add_para(
    doc,
    "二者在技术构思及最终实现的功能上存在根本性差异。具体而言：",
)
add_para(
    doc,
    "其一，应用场景与所连接对象不同。对比文件1用于膝关节置换，球面副连接的是“关节面骨连接件”与“单根髓内柄”，植入对象为同一骨骼的一端；本案用于长管状骨截骨矫形，球面副连接的是分别植入两个骨段的第一假体与第二假体，植入对象为截骨后彼此分离的近端骨段与远端骨段。",
)
add_para(
    doc,
    "其二，结构拓扑不同。对比文件1为“骨连接件＋单髓内柄＋固定部件”的单侧悬臂结构，髓内柄仅向一侧延伸进入单一髓腔；本案为“第一假体＋第二假体”的对向双柄结构，两髓内柄自球面副向相反两侧延伸，分别进入两个骨段髓腔，球面副位于两截骨端之间。",
)
add_para(
    doc,
    "其三，球面副所调节的对象不同。对比文件1调节的是髓内柄相对关节面假体的安装倾角，关节面方位已由置换截骨面预先确定，球面副并不改变两段骨骼之间的相对力线；本案调节的是截骨后两个骨段之间的相对成角与力线，球面副本身即为矫形的几何枢纽。",
)
add_para(
    doc,
    "其四，最终实现的功能不同。对比文件1实现的是置换假体与患者髓腔形态的匹配，降低延长杆皮质撞击风险；本案实现的是截骨矫形精度的提高以及矫正后力线在愈合期内的髓内维持。对比文件1未公开第一假体与第二假体、未公开两髓内柄自球面副向相反两侧延伸并分别插入两个骨段髓腔以调整、维持力线的技术特征。",
)
add_para(
    doc,
    "本案相对于对比文件1具备新颖性和创造性。",
    bold=True,
)

add_section(doc, "三、对比文件2的分析")
add_subsection(doc, "（一）对比文件2公开的技术方案")
add_para(
    doc,
    "对比文件2（CN209253032U）公开一种用于胫骨近端高位截骨的矫正装置，属于手术导板/导向器械。该装置包括与胫骨近端内侧骨质表面贴合的贴合导板，贴合导板上设有横向截骨槽、轨道矫正机构和导板固定装置；轨道矫正机构包括位于横向截骨槽一侧的轨迹槽以及跟随截骨后部分胫骨沿轨迹槽偏摆的定位器；轨迹槽中设置至少两个矫正位，定位器可选择性地定位在任一矫正位。使用时，依据CT数据个性化设计并3D打印导板，固定于胫骨近端内侧后沿横向截骨槽截骨，再将矫正针沿轨迹槽由一个矫正孔移动至另一个矫正孔，完成预定夹角（通常为3°–6°）的一次性摆正。该装置为术中临时使用的导向、定位器械，术毕取出，并不作为植入物留置于体内。",
)

add_subsection(doc, "（二）与本案的技术构思及功能对比")
add_para(
    doc,
    "对比文件2采用“贴合导板＋轨迹槽离散矫正位”的预定路径式术中导向：预先依据CT扫描在贴合导板上划定横向截骨槽、轨迹槽及至少两个矫正孔，截骨后检查定位器是否到达下一矫正位，未到位则沿轨迹槽继续偏摆，到达预设矫正孔后插入套筒锁定。其矫正路径、矫正角度在术前即被导板几何所固定，术中只能在既定轨迹上作单平面、离散位点的摆正，装置本身不植入、不跨越截骨线建立永久髓内连接，也不承担术后力线的长期维持。",
)
add_para(
    doc,
    "本专利采用“植入式双髓内柄＋球面副连续全向锁定”的体内矫形维持：第一假体、第二假体分别插入两骨段髓腔并经球面副相连，术中可绕球心在多个解剖学平面内连续、精细地试摆力线，到位后以顶紧螺钉等手段锁定球面副，两髓内柄在髓腔内形成跨越截骨线的内在支撑，于愈合期维持已矫正的力线。",
)
add_para(
    doc,
    "二者在技术构思及最终实现的功能上存在根本性差异。具体而言：",
)
add_para(
    doc,
    "其一，产品属性不同。对比文件2是一次性手术导板器械，功能止于术中截骨导向与预定角度摆正；本案是可植入的骨假体，功能延伸至术中全向调角、锁定以及术后髓内支撑与力线维持。",
)
add_para(
    doc,
    "其二，矫形自由度不同。对比文件2的轨迹槽相对横向截骨槽垂直设置，两矫正孔轴线夹角等于术前设定的待矫正角，属于单一平面内、在有限个矫正位之间的离散跳变；本案球面副允许外球面相对内球面绕球心作多方向连续转动，可针对同时具有多方向分量的成角畸形进行全向、精细微调，并不依赖术前固化的单一轨迹。",
)
add_para(
    doc,
    "其三，力线维持方式不同。对比文件2矫正完成后导板即失去固定意义，两骨段后续仍需依赖其他固定方式；本案两髓内柄分别插入两骨段髓腔，经球面副将两骨段连接为一体，在截骨线两侧建立起髓内载荷传递路径。对比文件2未公开第一假体、第二假体、外球面与内球面构成的球面副，亦未公开两髓内柄对向插入两骨段髓腔以调整、维持力线的任何技术特征。",
)
add_para(
    doc,
    "本案相对于对比文件2具备新颖性和创造性。",
    bold=True,
)

add_section(doc, "四、对比文件1与对比文件2结合的分析")
add_para(
    doc,
    "对比文件1采用“单柄—关节面球面适配”的被动匹配式调节，解决的是膝关节置换中髓内柄与髓腔的匹配问题；对比文件2采用“贴合导板＋轨迹槽离散矫正位”的预定路径式术中导向，解决的是胫骨近端高位截骨术中截骨位置与单一平面预定角度的精确摆正问题。二者分属膝关节置换假体与截骨导板器械两条技术路线，所针对的临床场景、产品形态及功能终点均不相同。",
)
add_para(
    doc,
    "即便将对比文件1的球面副调节手段与对比文件2的截骨矫形应用场景简单拼合，对比文件1给出的仍是“关节面假体＋单根髓内柄”的置换装配结构，对比文件2给出的仍是“导板轨迹离散摆正”的一次性手术导向，二者均未给出将第一假体、第二假体经球面副相互连接、两髓内柄自球面副向相反两侧延伸并分别插入截骨后两个骨段髓腔以调整并维持力线的技术启示。本领域技术人员没有动机将对比文件1中的骨连接件（股骨髁假体或胫骨托）改造成带有第二髓内柄的第二假体，也没有动机将对比文件2中的临时贴合导板替换为可植入、可连续全向锁定的双柄骨假体。上述结合既不能破坏本案独立权利要求的新颖性，亦不能使本案相对于现有技术的区别技术特征对本领域技术人员而言变得显而易见。",
)
add_para(
    doc,
    "本案相对于对比文件1、对比文件2及其结合仍具备新颖性和创造性。",
    bold=True,
)

add_section(doc, "五、从属权利要求的简要评述")
add_para(
    doc,
    "在独立权利要求具备新颖性、创造性的前提下，从属权利要求2–10所限定的附加技术特征（内球面包绕角大于180°及防脱唇、顶紧螺钉锁定球面副、髓内柄轴向凸棱/凹槽及锁定孔、非圆形横截面防转、空心多孔结构及骨诱导材料填充、羟基磷灰石涂层、连接部支承面与外径适配等）对比文件1、2均未公开，亦非本领域解决“截骨后两骨段全向调角并髓内维持力线”这一技术问题时的常规选择。上述从属权利要求同样具备新颖性和创造性。",
)

add_section(doc, "六、对比要点一览")

table = doc.add_table(rows=6, cols=4)
table.style = "Table Grid"
headers = ["对比维度", "对比文件1\nCN116965979A", "对比文件2\nCN209253032U", "本案\nPA26-0865-U-A"]
for i, h in enumerate(headers):
    set_cell_text(table.rows[0].cells[i], h, bold=True, center=True, size=10)
    set_cell_shading(table.rows[0].cells[i], "1F4E79")
    for p in table.rows[0].cells[i].paragraphs:
        for r in p.runs:
            r.font.color.rgb = RGBColor(255, 255, 255)

rows_data = [
    [
        "技术领域\n与产品属性",
        "膝关节置换用腿骨假体组件（植入物）",
        "胫骨近端高位截骨用矫正导板（手术器械）",
        "长管状骨截骨矫形用骨假体（植入物）",
    ],
    [
        "核心结构",
        "骨连接件（股骨髁/胫骨托）＋单根髓内柄＋固定部件；柄与骨连接件、柄与固定部件之间为球面副",
        "贴合导板＋横向截骨槽＋轨迹槽＋定位器；轨迹槽两端为离散矫正孔",
        "第一假体＋第二假体；两连接部构成球面副；两髓内柄自球面副向相反两侧延伸",
    ],
    [
        "调节对象\n与自由度",
        "单根髓内柄相对已确定关节面的安装倾角；万向转动，但只服务同一骨骼一端",
        "截骨后部分胫骨沿预定轨迹偏摆；单平面、两（或多）个离散矫正位",
        "截骨后两个骨段之间的相对力线；绕球心全向、连续、精细调整",
    ],
    [
        "力线维持",
        "不连接两个骨段，不承担截骨矫形后的力线维持",
        "导板术毕取出，不形成跨截骨线的髓内支撑",
        "两柄分别插入两骨段髓腔，经球面副将两骨段连为一体，术后维持力线",
    ],
    [
        "技术构思\n定性",
        "“单柄—关节面球面适配”的被动匹配式调节",
        "“贴合导板＋轨迹槽离散矫正位”的预定路径式术中导向",
        "“双柄对向植入＋球面副枢纽全向矫形”的主动规划式调节",
    ],
]
for r_i, row in enumerate(rows_data, start=1):
    fill = "D6E3F0" if r_i % 2 == 1 else "FFFFFF"
    for c_i, val in enumerate(row):
        set_cell_text(table.rows[r_i].cells[c_i], val, bold=(c_i == 0), size=9)
        set_cell_shading(table.rows[r_i].cells[c_i], fill)

# set column widths
widths = [Cm(3.0), Cm(4.2), Cm(4.2), Cm(4.4)]
for row in table.rows:
    for idx, cell in enumerate(row.cells):
        cell.width = widths[idx]

add_section(doc, "七、结论")
add_para(
    doc,
    "对比文件1采用“单柄—关节面球面适配”的被动匹配式调节，对比文件2采用“贴合导板＋轨迹槽离散矫正位”的预定路径式术中导向；本专利采用“双柄对向植入＋球面副枢纽全向矫形”的主动规划式调节。本案与两份对比文件在技术构思及最终实现的功能上均存在根本性差异。本案独立权利要求1及从属权利要求2–10相对于对比文件1、对比文件2以及二者的结合，具备《专利法》第二十二条第二款、第三款所规定的新颖性和创造性。",
    bold=True,
)

add_para(
    doc,
    "（本文件仅就用户指定的两份对比文件撰写检索报告之“分析部分”，供内部讨论与正式检索报告撰写使用。）",
    first_line=False,
    size=10.5,
    space_after=0,
)

out_path = "/workspace/PA26-0865-U-A_检索报告_分析部分.docx"
doc.save(out_path)
print("saved", out_path)
